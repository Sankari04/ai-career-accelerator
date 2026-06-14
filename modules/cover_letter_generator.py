"""
cover_letter_generator.py
AI Cover Letter Generator — persists generated letter in session state
so PDF and DOCX downloads work independently without re-generation.
"""

import io
import streamlit as st
from fpdf import FPDF
from docx import Document
from utils.api_client import generate_response
from utils import session_state as ss


# ── Export helpers ─────────────────────────────────────────────────────────────

class _PDF(FPDF):
    def header(self):
        self.set_font("Arial", "B", 15)
        self.cell(0, 10, "Professional Cover Letter", 0, 1, "C")
        self.ln(4)


def _build_pdf(content: str) -> bytes:
    pdf = _PDF()
    pdf.add_page()
    pdf.set_font("Arial", size=11)
    pdf.set_auto_page_break(auto=True, margin=15)

    for line in content.split("\n"):
        line = line.strip().replace("**", "")
        if not line:
            pdf.ln(4)
        else:
            line = line.encode("latin-1", "replace").decode("latin-1")
            pdf.multi_cell(0, 6, line)

    return pdf.output(dest="S").encode("latin-1")


def _build_docx(content: str) -> bytes:
    doc = Document()
    for line in content.split("\n"):
        line = line.strip().replace("**", "")
        doc.add_paragraph(line if line else "")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ── Prompt builder ─────────────────────────────────────────────────────────────

def _build_prompt(name, job_role, company_name, hiring_manager,
                  skills, experience) -> str:
    greeting = f"Dear {hiring_manager}," if hiring_manager else "Dear Hiring Manager,"
    return f"""
You are an expert career coach and professional writer.
Write a formal, highly compelling, ready-to-send cover letter.

Rules:
- Formal, confident tone.
- Max 4 paragraphs; concise and engaging.
- No placeholders like [Insert Date] — the letter must be complete.
- Focus on how the applicant's skills bring value to {company_name}.
- Output plain text only — no markdown headings or bullet points.

Applicant Name: {name}
Target Role: {job_role}
Target Company: {company_name}
Greeting: {greeting}
Key Skills: {skills}
Experience: {experience}
"""


# ── Main render ────────────────────────────────────────────────────────────────

def render():
    st.markdown("<div class='animate-fade-in'>", unsafe_allow_html=True)

    c1, c2 = st.columns([3, 1])
    with c1:
        st.markdown("<h1 style='margin-bottom:0;'>✉️ Cover Letter Builder</h1>",
                    unsafe_allow_html=True)
        st.markdown("<p style='color:#94a3b8;font-size:1.1rem;'>Generate a "
                    "persuasive, personalised cover letter — saved for the whole session.</p>",
                    unsafe_allow_html=True)
    with c2:
        st.markdown("<div style='text-align:right;padding-top:20px;'>"
                    "<span class='premium-badge'>PDF & DOCX Export</span></div>",
                    unsafe_allow_html=True)

    # ── Input form ────────────────────────────────────────────────────────────
    with st.form("cover_letter_form"):
        st.markdown("<div class='glass-card'>", unsafe_allow_html=True)
        st.markdown("<h3 style='margin-top:0;color:#a78bfa;'>Job Details</h3>",
                    unsafe_allow_html=True)
        c1, c2 = st.columns(2)
        with c1:
            name     = st.text_input("Your Full Name *", placeholder="Jane Doe")
            job_role = st.text_input("Target Job Role *",
                                     placeholder="Senior Product Manager")
        with c2:
            company_name    = st.text_input("Company Name *",
                                            placeholder="Google")
            hiring_manager  = st.text_input("Hiring Manager (optional)",
                                            placeholder="Leave blank if unknown")

        st.markdown("<h3 style='margin-top:10px;color:#a78bfa;'>Your Background</h3>",
                    unsafe_allow_html=True)
        skills     = st.text_area("Key Skills *",
                                  placeholder="Agile, SQL, Stakeholder Management",
                                  height=80)
        experience = st.text_area("Brief Experience Summary",
                                  placeholder="5 years leading cross-functional teams…",
                                  height=110)
        st.markdown("</div>", unsafe_allow_html=True)

        submitted = st.form_submit_button(
            "✨ Generate Cover Letter", use_container_width=True)

    # ── Generation ────────────────────────────────────────────────────────────
    if submitted:
        if not name or not job_role or not company_name or not skills:
            st.error("⚠️ Name, Job Role, Company Name, and Skills are required.")
        else:
            with st.spinner("🧠 Drafting your cover letter…"):
                prompt     = _build_prompt(name, job_role, company_name,
                                           hiring_manager, skills, experience)
                content    = generate_response(prompt)
                pdf_bytes  = _build_pdf(content)
                docx_bytes = _build_docx(content)

            ss.set_value("cover", "content",  content)
            ss.set_value("cover", "pdf",      pdf_bytes)
            ss.set_value("cover", "docx",     docx_bytes)
            ss.set_value("cover", "filename", name.replace(" ", "_"))
            st.success("✅ Cover letter generated and saved to your session!")

    # ── Display stored result ─────────────────────────────────────────────────
    _show_result()
    st.markdown("</div>", unsafe_allow_html=True)


def _show_result():
    data = ss.get("cover")
    if not data.get("content"):
        return

    st.markdown("<div class='glass-card animate-fade-in'>", unsafe_allow_html=True)
    st.markdown("### 👁️ Preview")
    st.markdown(
        "<div style='background:rgba(15,23,42,0.8);padding:28px;"
        "border-radius:12px;border:1px solid rgba(255,255,255,0.1);"
        "white-space:pre-wrap;'>",
        unsafe_allow_html=True)
    st.markdown(data["content"])
    st.markdown("</div></div>", unsafe_allow_html=True)

    st.markdown("<div class='glass-card animate-fade-in'>", unsafe_allow_html=True)
    st.markdown("### 📥 Download")
    fname = data.get("filename", "cover_letter")
    c1, c2 = st.columns(2)
    with c1:
        st.download_button("📄 Download PDF",
            data=data["pdf"],
            file_name=f"{fname}_Cover_Letter.pdf",
            mime="application/pdf",
            use_container_width=True,
            key="cover_pdf_dl")
    with c2:
        st.download_button("📝 Download DOCX",
            data=data["docx"],
            file_name=f"{fname}_Cover_Letter.docx",
            mime="application/vnd.openxmlformats-officedocument"
                 ".wordprocessingml.document",
            use_container_width=True,
            key="cover_docx_dl")
    st.markdown("</div>", unsafe_allow_html=True)
